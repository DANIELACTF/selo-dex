"""Testa o script único `moraex.py` — o encadeamento das etapas, o relatório
de carência e a validação das skills antes do upload."""
import json
from pathlib import Path

import openpyxl
import pytest

import moraex
from onboarding.alimentar_carteira import COL_COMPETENCIA, situacao_carencia
from onboarding.pastas import gerar as gerar_pastas
from onboarding.pastas import nome_pasta
from onboarding.skills_pack import LIMITE_DESCRICAO, processar, validar

from tests.test_implantacao import _carteira_fake, _csv_pendentes, _preencher

EMAIL = Path("fixtures/exemplo_thays_2026-08-20.txt")

FRONTMATTER_OK = """---
name: minha-skill
description: 'Faz alguma coisa: e cita dois pontos no meio.'
---

# Corpo
"""


def _skill(tmp_path: Path, nome: str, conteudo: str) -> Path:
    pasta = tmp_path / "skills" / nome
    pasta.mkdir(parents=True)
    (pasta / "SKILL.md").write_text(conteudo, encoding="utf-8")
    return pasta


def _planilha(tmp_path: Path) -> Path:
    from onboarding.planilha_particularidades import gerar

    return gerar(_csv_pendentes(tmp_path), tmp_path / "particularidades.xlsx", "08/2026")


# --------------------------------------------------------------- etapa 1

def test_triagem_gera_fichas_e_lista_de_pendentes(tmp_path, capsys):
    pendentes = tmp_path / "pendentes.csv"
    codigo = moraex.main([
        "triagem", "--email", str(EMAIL), "--sem-consulta",
        "--fichas", str(tmp_path / "fichas"), "--pendentes", str(pendentes),
    ])
    assert codigo == 0
    assert pendentes.exists()
    assert list((tmp_path / "fichas").glob("*.pdf"))
    assert "Próximo passo" in capsys.readouterr().out


def test_triagem_em_json_traz_alertas_de_certificado(tmp_path, capsys):
    moraex.main([
        "triagem", "--email", str(EMAIL), "--sem-consulta", "--json",
        "--fichas", str(tmp_path / "fichas"), "--pendentes", str(tmp_path / "p.csv"),
    ])
    dados = json.loads(capsys.readouterr().out)
    assert dados["comando"] == "triagem"
    assert len(dados["empresas"]) == 5
    assert any("GABRIELE" in a for a in dados["alertas_certificado"])


def test_etapa1_encadeia_triagem_e_planilha(tmp_path):
    saida = tmp_path / "particularidades.xlsx"
    codigo = moraex.main([
        "etapa1", "--email", str(EMAIL), "--sem-consulta",
        "--fichas", str(tmp_path / "fichas"), "--pendentes", str(tmp_path / "p.csv"),
        "--saida", str(saida), "--competencia", "08/2026",
    ])
    assert codigo == 0
    ws = openpyxl.load_workbook(saida)["Particularidades"]
    assert ws.max_row == 6  # cabeçalho + 5 empresas do lote


def test_planilha_sem_triagem_previa_falha_com_instrucao(tmp_path, capsys):
    codigo = moraex.main(["planilha", "--pendentes", str(tmp_path / "nao-existe.csv")])
    assert codigo == 1
    assert "moraex.py triagem" in capsys.readouterr().out


# --------------------------------------------------------------- pastas

@pytest.mark.parametrize("razao,esperado", [
    ("C LORENA DISTRIBUIDORA LTDA", "C LORENA DISTRIBUIDORA"),
    ("GRÁFICA SQUARE LTDA", "GRAFICA SQUARE"),
    ("INJECT PHARMA MANIPULAÇÃO EIRELI", "INJECT PHARMA MANIPULACAO"),
])
def test_nome_da_pasta_sem_acento_e_sem_sufixo_societario(razao, esperado):
    assert nome_pasta(razao) == esperado


def test_pastas_gera_csv_do_lote_e_comando(tmp_path):
    r = gerar_pastas(_planilha(tmp_path), tmp_path / "lote.csv", raiz="F:\\Clientes")
    linhas = (tmp_path / "lote.csv").read_text(encoding="utf-8-sig").splitlines()
    assert linhas[0] == "Numero,Nome"
    assert "1091,C LORENA DISTRIBUIDORA" in linhas
    assert "F:\\Clientes" in r["comando"] and "lote.csv" in r["comando"]
    assert r["pastas"] == ["1091 - C LORENA DISTRIBUIDORA", "1092 - GRAFICA SQUARE"]


# --------------------------------------------------------------- etapa 2

def test_etapa2_gera_csv_e_carteira_nova(tmp_path):
    planilha = _planilha(tmp_path)
    carteira = _carteira_fake(tmp_path)
    saida = tmp_path / "carteira-nova.xlsx"

    codigo = moraex.main([
        "etapa2", "--planilha", str(planilha), "--carteira", str(carteira),
        "--saida-carteira", str(saida), "--saida-csv", str(tmp_path / "lote.csv"),
        "--competencia", "08/2026",
    ])
    assert codigo == 0
    assert (tmp_path / "lote.csv").exists()
    assert saida.exists()
    # as duas empresas entraram em carência, não na carteira do analista
    wb = openpyxl.load_workbook(saida)
    assert wb["Pendentes Daniela"].max_row == 3
    assert wb["Carteira Completa"].max_row == 2


# --------------------------------------------------------------- status

def test_status_separa_liberadas_de_quem_ainda_espera(tmp_path):
    carteira = _carteira_fake(tmp_path)
    wb = openpyxl.load_workbook(carteira)
    wsp = wb["Pendentes Daniela"]
    wsp.cell(row=1, column=9, value=COL_COMPETENCIA)
    wsp.append(["1091", "PRONTA LTDA", "11.111.111/0001-11", "", "", "Dulce Neves", "", "", "05/2026"])
    wsp.append(["1092", "ESPERANDO LTDA", "22.222.222/0001-22", "", "", "", "", "", "08/2026"])
    wsp.append(["999", "ANTIGA LTDA", "33.333.333/0001-33", "", "", "", "", "", ""])
    wb.save(carteira)

    r = situacao_carencia(carteira, "08/2026")
    assert [i["numero"] for i in r["liberadas"]] == ["1091"]
    assert [i["numero"] for i in r["em_carencia"]] == ["1092"]
    assert [i["numero"] for i in r["sem_competencia"]] == ["999"]
    assert r["em_carencia"][0]["libera_em"] == "11/2026"
    assert r["em_carencia"][0]["faltam"] == 3
    assert r["total"] == 3


def test_status_nao_escreve_na_carteira(tmp_path):
    carteira = _carteira_fake(tmp_path)
    antes = carteira.read_bytes()
    moraex.main(["status", "--carteira", str(carteira), "--competencia", "08/2026"])
    assert carteira.read_bytes() == antes


# --------------------------------------------------------------- skills

def test_skills_do_repositorio_passam_na_validacao():
    for r in processar(Path(".claude/skills")):
        assert r.ok, f"{r.pasta.name}: {r.erros}"
        assert r.tamanho_descricao <= LIMITE_DESCRICAO


def test_description_acima_do_limite_e_recusada(tmp_path):
    longa = "x" * (LIMITE_DESCRICAO + 1)
    pasta = _skill(tmp_path, "minha-skill", f"---\nname: minha-skill\ndescription: '{longa}'\n---\n")
    assert any("limite do uploader" in e for e in validar(pasta).erros)


def test_falta_de_frontmatter_e_recusada(tmp_path):
    pasta = _skill(tmp_path, "minha-skill", "# Skill sem frontmatter\n")
    assert any("frontmatter" in e for e in validar(pasta).erros)


def test_dois_pontos_sem_aspas_e_recusado(tmp_path):
    pasta = _skill(
        tmp_path, "minha-skill",
        "---\nname: minha-skill\ndescription: Faz algo: e quebra o YAML\n---\n",
    )
    assert any("aspas" in e for e in validar(pasta).erros)


def test_dois_pontos_entre_aspas_passa(tmp_path):
    pasta = _skill(tmp_path, "minha-skill", FRONTMATTER_OK)
    assert validar(pasta).ok


def test_name_precisa_bater_com_a_pasta(tmp_path):
    pasta = _skill(tmp_path, "outra-pasta", FRONTMATTER_OK)
    assert any("não bate com o nome da pasta" in e for e in validar(pasta).erros)


def test_zip_sai_com_skill_md_na_raiz(tmp_path):
    import zipfile

    _skill(tmp_path, "minha-skill", FRONTMATTER_OK)
    (tmp_path / "skills" / "minha-skill" / "modelo.html").write_text("<p>ok</p>", encoding="utf-8")

    r = processar(tmp_path / "skills", tmp_path / "dist")[0]
    assert r.ok and r.zip is not None
    with zipfile.ZipFile(r.zip) as z:
        assert sorted(z.namelist()) == ["SKILL.md", "modelo.html"]
